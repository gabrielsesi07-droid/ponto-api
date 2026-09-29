"""Read-only corporate document ingestion. Output must stay in an ignored/private folder.

Usage: python scripts/extract-library.py --source checklist PATH --source catalog PATH
       --source manual PATH --output work/library
No executable is opened or run. ZIP paths are logical provenance, never extraction paths.
"""
import argparse
import hashlib
import io
import json
import re
import unicodedata
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

SUPPORTED = {'.pdf', '.docx', '.xlsx', '.txt', '.doc'}
MAX_FILE = 40 * 1024 * 1024
MAX_ARCHIVE_BYTES = 300 * 1024 * 1024


def normalize(text):
    return ''.join(c for c in unicodedata.normalize('NFKD', text.lower()) if not unicodedata.combining(c))


def extract(data, extension):
    sections, warnings = [], []
    def add(locator, value):
        text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value)).strip()
        for offset in range(0, len(text), 10000):
            sections.append({'locator': locator, 'text': text[offset:offset + 10000]})
    if extension == '.pdf':
        reader = PdfReader(io.BytesIO(data))
        empty = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ''
            if len(text.strip()) < 30:
                empty.append(i + 1)
            add(f'Página {i + 1}', text)
        if empty:
            warnings.append('Páginas sem texto suficiente; imagens não passaram por OCR: ' + ', '.join(map(str, empty)))
    elif extension == '.docx':
        doc = Document(io.BytesIO(data))
        # Preserve paragraph/table sequence, not an artificial concatenation of all tables.
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        for i, element in enumerate(doc.element.body):
            if element.tag.endswith('}p'):
                add(f'Bloco {i + 1}', Paragraph(element, doc).text)
            elif element.tag.endswith('}tbl'):
                for j, row in enumerate(Table(element, doc).rows):
                    add(f'Tabela {i + 1}, linha {j + 1}', ' | '.join(cell.text for cell in row.cells))
        for i, section in enumerate(doc.sections):
            add(f'Cabeçalho {i + 1}', '\n'.join(p.text for p in section.header.paragraphs))
            add(f'Rodapé {i + 1}', '\n'.join(p.text for p in section.footer.paragraphs))
        if doc.inline_shapes:
            warnings.append('Contém imagens; consulte o original para desenhos, fotos e campos visuais.')
    elif extension == '.xlsx':
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        for sheet in book:
            if (sheet.max_row or 0) > 100000 or (sheet.max_column or 0) > 1000:
                warnings.append(f'Aba {sheet.title}: dimensões excessivas; leitura não realizada.')
                continue
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if any(value is not None for value in row):
                    add(f'Aba {sheet.title}, linha {i + 1}', ' | '.join('' if v is None else str(v) for v in row))
        book.close()
        warnings.append('Valores salvos da planilha; fórmulas não recalculadas. Consulte o original para layout e gráficos.')
    elif extension == '.txt':
        try:
            value = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            value = data.decode('cp1252', errors='replace')
        add('Texto', value)
    else:
        warnings.append('Formato Word antigo (.doc): original preservado, sem extração de texto.')
    if not sections:
        warnings.append('Sem texto pesquisável. Necessita leitura do original ou OCR.')
    return sections, warnings


def run(sources, output):
    output.mkdir(parents=True, exist_ok=True)
    originals = output / 'originals'
    originals.mkdir(exist_ok=True)
    documents, skipped, archives, errors = {}, [], [], []
    def accept(data, source, category, depth=0):
        name = source.replace('\\', '/').split('/')[-1]
        extension = Path(name).suffix.lower()
        if name.startswith('~$'):
            skipped.append({'source': source, 'reason': 'arquivo temporário'})
            return
        if extension == '.zip':
            if depth >= 4:
                errors.append({'source': source, 'reason': 'limite de profundidade de ZIP'})
                return
            try:
                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    entries = archive.infolist()
                    if len(entries) > 5000 or sum(e.file_size for e in entries) > MAX_ARCHIVE_BYTES:
                        raise ValueError('ZIP excede limite de expansão')
                    archives.append(source)
                    for entry in entries:
                        if entry.is_dir():
                            continue
                        child = source + '!/' + entry.filename
                        if entry.file_size > MAX_FILE or (entry.compress_size and entry.file_size / entry.compress_size > 1000):
                            skipped.append({'source': child, 'reason': 'limite de tamanho/expansão'})
                        elif Path(entry.filename).suffix.lower() in SUPPORTED | {'.zip'}:
                            accept(archive.read(entry), child, category, depth + 1)
                        else:
                            skipped.append({'source': child, 'reason': 'software ou formato não documental'})
            except Exception as error:
                errors.append({'source': source, 'reason': type(error).__name__ + ': ' + str(error)[:200]})
            return
        if extension not in SUPPORTED:
            skipped.append({'source': source, 'reason': 'software ou formato não documental'})
            return
        if len(data) > MAX_FILE:
            skipped.append({'source': source, 'reason': 'documento excede 40 MB'})
            return
        digest = hashlib.sha256(data).hexdigest()
        obsolete = any(x in normalize(source) for x in ('obsoleto', 'nao utilizar'))
        origin = {'path': source, 'category': category, 'obsolete': obsolete}
        if digest in documents:
            documents[digest]['origins'].append(origin)
            # Conservatively quarantine any copy marked obsolete until explicitly reviewed.
            documents[digest]['obsolete'] |= obsolete
            return
        try:
            sections, warnings = extract(data, extension)
        except Exception as error:
            sections, warnings = [], ['Falha na extração: ' + type(error).__name__ + '. Original preservado.']
        (originals / (digest + extension)).write_bytes(data)
        documents[digest] = dict(hash=digest, name=name, extension=extension, size=len(data),
            category=category, obsolete=obsolete, origins=[origin], sections=sections, warnings=warnings)
    for category, raw_path in sources:
        root = Path(raw_path)
        if not root.is_dir():
            raise FileNotFoundError(f'Fonte inacessível: {root}')
        for path in sorted(root.rglob('*')):
            if not path.is_file():
                continue
            try:
                if path.suffix.lower() not in SUPPORTED | {'.zip'} or path.name.startswith('~$'):
                    skipped.append({'source': str(path), 'reason': 'software, temporário ou formato não documental'})
                elif path.stat().st_size > MAX_ARCHIVE_BYTES:
                    skipped.append({'source': str(path), 'reason': 'limite de tamanho'})
                else:
                    accept(path.read_bytes(), str(path), category)
            except OSError as error:
                errors.append({'source': str(path), 'reason': str(error)})
    result = dict(created_at=datetime.now(timezone.utc).isoformat(), documents=list(documents.values()),
                  skipped=skipped, archives=archives, errors=errors)
    (output / 'manifest.json').write_text(json.dumps(result, ensure_ascii=False), encoding='utf-8')
    print(json.dumps(dict(documents=len(documents), origins=sum(len(d['origins']) for d in documents.values()),
        bytes=sum(d['size'] for d in documents.values()), sections=sum(len(d['sections']) for d in documents.values()),
        obsolete=sum(d['obsolete'] for d in documents.values()), no_text=sum(not d['sections'] for d in documents.values()),
        skipped=len(skipped), archives=len(archives), errors=len(errors)), ensure_ascii=False))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', nargs=2, action='append', required=True, metavar=('CATEGORY', 'PATH'))
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    if any(category not in ('checklist', 'catalog', 'manual') for category, _ in args.source):
        parser.error('Categorias aceitas: checklist, catalog, manual')
    run(args.source, Path(args.output))
